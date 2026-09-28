import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn
import os, shutil

# Paths
base_dir = r"c:\Users\user\Documents\Fisiomedi"
web_docs_dir = r"c:\Users\user\Documents\Fisiomedi\web\docs"
diag_dir = os.path.join(web_docs_dir, "generated_diagrams")
old_img_dir = os.path.join(base_dir, "docs", "extracted_images")
cibertec_logo = os.path.join(old_img_dir, "image_28.jpeg")

output_docx_web = os.path.join(web_docs_dir, "Proyecto_Registro_Citas_FISIOMEDI_Web_EFSRT_II.docx")
output_docx_root = os.path.join(base_dir, "docs", "Proyecto_Registro_Citas_FISIOMEDI_Web_EFSRT_II.docx")

doc = docx.Document()

# Page setup: Margins
for section in doc.sections:
    section.top_margin = Inches(1.0)
    section.bottom_margin = Inches(1.0)
    section.left_margin = Inches(1.0)
    section.right_margin = Inches(1.0)
    
    # Header & Footer
    header = section.header
    hp = header.paragraphs[0]
    hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    hrun = hp.add_run("FISIOMEDI | Sistema Web de Gestión de Citas y Resultados Médicos — EFSRT II")
    hrun.font.size = Pt(8.5)
    hrun.font.color.rgb = RGBColor(100, 116, 139)
    
    footer = section.footer
    fp = footer.paragraphs[0]
    fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    frun = fp.add_run("Instituto Superior Tecnológico CIBERTEC · 2026")
    frun.font.size = Pt(8.5)
    frun.font.color.rgb = RGBColor(148, 163, 184)

# Color constants
COLOR_NAVY = RGBColor(30, 58, 138)    # #1e3a8a
COLOR_BLUE = RGBColor(37, 99, 235)    # #2563eb
COLOR_DARK = RGBColor(15, 23, 42)     # #0f172a
COLOR_MUTED = RGBColor(100, 116, 139) # #64748b

def set_cell_background(cell, hex_color):
    shading_xml = f'<w:shd {nsdecls("w")} w:fill="{hex_color}"/>'
    cell._tc.get_or_add_tcPr().append(parse_xml(shading_xml))

def add_title(text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(12)
    run = p.add_run(text)
    run.font.name = "Arial"
    run.font.size = Pt(18)
    run.font.bold = True
    run.font.color.rgb = COLOR_NAVY
    return p

def add_heading_1(text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(18)
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.keep_with_next = True
    run = p.add_run(text)
    run.font.name = "Arial"
    run.font.size = Pt(13.5)
    run.font.bold = True
    run.font.color.rgb = COLOR_NAVY
    return p

def add_heading_2(text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(12)
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.keep_with_next = True
    run = p.add_run(text)
    run.font.name = "Arial"
    run.font.size = Pt(11.5)
    run.font.bold = True
    run.font.color.rgb = COLOR_BLUE
    return p

def add_heading_3(text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(8)
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.keep_with_next = True
    run = p.add_run(text)
    run.font.name = "Arial"
    run.font.size = Pt(10.5)
    run.font.bold = True
    run.font.color.rgb = COLOR_DARK
    return p

def add_p(text, bold_prefix="", italic=False):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.line_spacing = 1.15
    if bold_prefix:
        r_pre = p.add_run(bold_prefix)
        r_pre.font.name = "Arial"
        r_pre.font.size = Pt(10)
        r_pre.font.bold = True
        r_pre.font.color.rgb = COLOR_DARK
    r_body = p.add_run(text)
    r_body.font.name = "Arial"
    r_body.font.size = Pt(10)
    r_body.font.italic = italic
    r_body.font.color.rgb = COLOR_DARK
    return p

def add_bullet(bold_prefix, text):
    p = doc.add_paragraph(style="List Bullet")
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.space_after = Pt(3)
    p.paragraph_format.line_spacing = 1.15
    r_pre = p.add_run(bold_prefix)
    r_pre.font.name = "Arial"
    r_pre.font.size = Pt(10)
    r_pre.font.bold = True
    r_pre.font.color.rgb = COLOR_DARK
    
    r_body = p.add_run(" " + text)
    r_body.font.name = "Arial"
    r_body.font.size = Pt(10)
    r_body.font.color.rgb = COLOR_DARK
    return p

def add_image_figure(img_path, caption, width=Inches(6.2)):
    if os.path.exists(img_path):
        p_img = doc.add_paragraph()
        p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_img.paragraph_format.space_before = Pt(8)
        p_img.paragraph_format.space_after = Pt(4)
        run = p_img.add_run()
        run.add_picture(img_path, width=width)
        
        p_cap = doc.add_paragraph()
        p_cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_cap.paragraph_format.space_after = Pt(10)
        r_cap = p_cap.add_run(f"Figura: {caption}\nFuente: Elaboración propia del equipo de desarrollo, 2026.")
        r_cap.font.name = "Arial"
        r_cap.font.size = Pt(8.5)
        r_cap.font.italic = True
        r_cap.font.color.rgb = COLOR_MUTED

# ==================== CARÁTULA ====================
p_yr = doc.add_paragraph()
p_yr.alignment = WD_ALIGN_PARAGRAPH.CENTER
r_yr = p_yr.add_run("“Año de la Esperanza y el Fortalecimiento de la Democracia”")
r_yr.font.name = "Arial"
r_yr.font.size = Pt(10)
r_yr.font.italic = True
r_yr.font.color.rgb = COLOR_MUTED

# Cibertec Logo
if os.path.exists(cibertec_logo):
    p_logo = doc.add_paragraph()
    p_logo.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_logo.paragraph_format.space_before = Pt(12)
    p_logo.paragraph_format.space_after = Pt(12)
    p_logo.add_run().add_picture(cibertec_logo, width=Inches(2.2))

p_inst = doc.add_paragraph()
p_inst.alignment = WD_ALIGN_PARAGRAPH.CENTER
r_inst = p_inst.add_run("INSTITUTO SUPERIOR TECNOLÓGICO CIBERTEC\nESCUELA DE TECNOLOGÍAS DE LA INFORMACIÓN")
r_inst.font.name = "Arial"
r_inst.font.size = Pt(13)
r_inst.font.bold = True
r_inst.font.color.rgb = COLOR_NAVY

p_proj = doc.add_paragraph()
p_proj.alignment = WD_ALIGN_PARAGRAPH.CENTER
p_proj.paragraph_format.space_before = Pt(24)
p_proj.paragraph_format.space_after = Pt(6)
r_proj = p_proj.add_run("“SISTEMA WEB DE GESTIÓN Y RESERVA DE CITAS FISIOTERAPÉUTICAS CON PORTAL DE RESULTADOS MÉDICOS - FISIOMEDI”")
r_proj.font.name = "Arial"
r_proj.font.size = Pt(15)
r_proj.font.bold = True
r_proj.font.color.rgb = COLOR_NAVY

p_crs = doc.add_paragraph()
p_crs.alignment = WD_ALIGN_PARAGRAPH.CENTER
p_crs.paragraph_format.space_after = Pt(28)
r_crs = p_crs.add_run("PROYECTO DE EXPERIENCIA FORMATIVA EN SITUACIÓN REAL DE TRABAJO II (EFSRT II)")
r_crs.font.name = "Arial"
r_crs.font.size = Pt(11)
r_crs.font.bold = True
r_crs.font.color.rgb = COLOR_BLUE

# Info Box (Table)
info_tbl = doc.add_table(rows=6, cols=2)
info_tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
info_tbl.autofit = False

labels_data = [
    ("Profesor:", "JEAN CARLOS LAURENTE CHACON"),
    ("Ciclo / Aula / Semestre:", "Ciclo 3 · Aula 5590 · Semestre 2026 - 03"),
    ("Coordinador de Grupo:", "i202513002 - Kevin Usnayo Navarro"),
    ("Integrantes del Proyecto:", "• I202514348 - Jennifer Milagros Raquel Alarcón Calixto\n• i202513734 - Pedro Fernández Lores\n• i202514654 – Aldavic Jerymoth Zambrano Muñoz\n• i202513856 - Emilio Josué Solis Fernández"),
    ("Empresa Beneficiaria:", "FISIOMEDI - Centro Especializado de Fisioterapia y Rehabilitación"),
    ("Fecha de Presentación:", "Septiembre de 2026 · Lima, Perú")
]

for row_idx, (lbl, val) in enumerate(labels_data):
    row = info_tbl.rows[row_idx]
    c0 = row.cells[0]
    c1 = row.cells[1]
    c0.width = Inches(2.2)
    c1.width = Inches(4.2)
    
    p0 = c0.paragraphs[0]
    p0.paragraph_format.space_after = Pt(2)
    r0 = p0.add_run(lbl)
    r0.font.name = "Arial"
    r0.font.size = Pt(9.5)
    r0.font.bold = True
    r0.font.color.rgb = COLOR_NAVY
    
    p1 = c1.paragraphs[0]
    p1.paragraph_format.space_after = Pt(2)
    r1 = p1.add_run(val)
    r1.font.name = "Arial"
    r1.font.size = Pt(9.5)
    r1.font.color.rgb = COLOR_DARK

doc.add_page_break()

# ==================== 1. INTRODUCCIÓN ====================
add_heading_1("1. Introducción")
add_p(
    "El presente proyecto de ingeniería de software se desarrolla en el marco del curso de Experiencia Formativa en Situación Real de Trabajo II (EFSRT II) del Instituto Superior Tecnológico CIBERTEC. Su propósito central es diseñar, construir y desplegar una solución tecnológica moderna, profesional y altamente escalable que resuelva de manera integral las necesidades operativas de FISIOMEDI, una empresa especializada en servicios fisioterapéuticos y de rehabilitación integral."
)
add_p(
    "Tradicionalmente, muchas instituciones de salud y centros terapéuticos en etapa de crecimiento gestionan sus citas mediante cuadernos físicos, registros manuscritos o herramientas ofimáticas no centralizadas. Esta modalidad genera pérdida recurrente de información, errores por solapamiento de horarios entre terapeutas, falta de trazabilidad clínica y una notable incomodidad para los pacientes, quienes deben acudir físicamente o realizar múltiples llamadas para agendar turnos y recoger resultados de exámenes diagnósticos (como resonancias magnéticas, ecografías o radiografías)."
)
add_p(
    "Frente a este desafío, se ha desarrollado el Sistema Web FISIOMEDI utilizando una arquitectura moderna basada en Next.js 16 (React Server Components), TypeScript, base de datos relacional PostgreSQL en la nube (Neon Serverless) y estilos con Tailwind CSS. La plataforma no se limita únicamente a un formulario de citas; integra un portal de autoservicio para pacientes con autenticación segura por DNI, un panel administrativo en tiempo real con asignación dinámica de especialistas (médicos y fisioterapeutas), y un módulo de carga y publicación directa de resultados e imágenes médicas en formato digital."
)
add_p(
    "A través de este proyecto, el equipo de estudiantes aplica de forma rigurosa los estándares de la metodología RUP (Rational Unified Process) para el modelado del negocio, el paradigma de programación orientada a objetos en entornos web fullstack, el diseño de arquitecturas en capas y las mejores prácticas de seguridad informática en el tratamiento de historias clínicas y datos médicos."
)

# ==================== 2. JUSTIFICACIÓN ====================
add_heading_1("2. Justificación")
add_p(
    "La transformación digital del centro fisioterapéutico FISIOMEDI se justifica plenamente en tres dimensiones estratégicas: operativa, asistencial y tecnológica."
)
add_bullet(
    "Justificación Operativa:",
    "Elimina el uso de cuadernos y planillas de papel para el registro de citas. Automatiza la verificación de disponibilidad horaria por consultorio y especialista, permitiendo al personal de recepción agendar citas internas en segundos y evitando duplicidades de turnos."
)
add_bullet(
    "Justificación Asistencial y de Servicio al Paciente:",
    "Permite que los pacientes reserven sus sesiones las 24 horas del día desde cualquier dispositivo móvil o computadora. Además, resuelve un cuello de botella crítico: la entrega de resultados de estudios diagnósticos. Con el nuevo sistema, el fisioterapeuta sube directamente las imágenes o PDFs de resonancias y radiografías a la ficha del paciente, y este puede descargarlos en cualquier momento desde su cuenta privada."
)
add_bullet(
    "Justificación Tecnológica y Académica:",
    "Demuestra cómo una solución web fullstack construida con Next.js y PostgreSQL en la nube ofrece alta disponibilidad, costo de infraestructura optimizado mediante serverless (Neon y Vercel), y rigurosidad técnica al implementar modelos de datos relacionales íntegros con claves foráneas, índices de búsqueda acelerada y hashing criptográfico para contraseñas de usuarios."
)

# ==================== 3. BENEFICIARIOS ====================
add_heading_1("3. Beneficiarios del Proyecto")

add_heading_2("3.1. Beneficiarios Directos")
add_bullet(
    "Personal Administrativo y de Recepción:",
    "Cuentan con un panel centralizado (/admin/citas) para visualizar reservas en tiempo real, filtrar por especialista, validar o reprogramar turnos y crear expedientes clínicos digitales de manera ordenada."
)
add_bullet(
    "Fisioterapeutas y Médicos Especialistas:",
    "Acceden a su agenda de atenciones diarias sin cruces de horarios. Al concluir una sesión o recibir un estudio médico (resonancia magnética, rayos X), pueden cargar de inmediato el archivo diagnóstico, registrar sus conclusiones y cerrar la cita automáticamente."
)
add_bullet(
    "Pacientes del Centro FISIOMEDI:",
    "Disponen de un portal personalizado (/mi-cuenta) para consultar próximas citas, revisar su historial terapéutico, gestionar sus datos de contacto y descargar sus informes y placas radiográficas en formato digital PDF sin trasladarse físicamente a la clínica."
)

add_heading_2("3.2. Beneficiarios Indirectos")
add_bullet(
    "Nuevos Clientes y Comunidad:",
    "Acceden a un sitio web profesional donde pueden conocer los servicios ofrecidos, el perfil del equipo terapéutico y agendar una cita en línea de forma ágil y transparente."
)
add_bullet(
    "Familiares y Cuidadores de Pacientes:",
    "Obtienen certeza en la programación de sesiones y rapidez al compartir resultados médicos con otros especialistas de la salud."
)
add_bullet(
    "Dirección y Gerencia de FISIOMEDI:",
    "Dispone de métricas confiables de atenciones por fisioterapeuta, estados de citas (confirmadas, pendientes, canceladas) e historial consolidado de pacientes para la toma de decisiones estratégicas."
)

# ==================== 4. OBJETIVOS DEL PROYECTO ====================
add_heading_1("4. Objetivos del Proyecto")

add_heading_2("4.1. Objetivo General")
add_p(
    "Desarrollar, modelar e implementar una plataforma web integral de gestión y reserva de citas fisioterapéuticas con módulo de resultados médicos digitales para FISIOMEDI, optimizando la productividad operativa del personal de salud y mejorando la experiencia de atención del paciente mediante una arquitectura web moderna, segura y disponible 24/7."
)

add_heading_2("4.2. Objetivos Específicos (SMART)")
add_bullet(
    "Objetivo SMART 1 (Digitalización de Citas):",
    "Digitalizar el 100% del proceso de agendamiento y reserva de citas de FISIOMEDI antes del cierre del semestre académico 2026-03, sustituyendo completamente los registros manuales en papel por almacenamiento relacional en PostgreSQL."
)
add_bullet(
    "Objetivo SMART 2 (Módulo de Resultados Diagnósticos):",
    "Implementar un módulo clínico que permita a los fisioterapeutas subir archivos de hasta 10 MB (PDF, imágenes de resonancia magnética, rayos X y ecografías), reduciendo a 0 minutos el tiempo de espera del paciente para acceder a sus resultados mediante el portal 'Mi Cuenta'."
)
add_bullet(
    "Objetivo SMART 3 (Asignación Inteligente de Especialistas):",
    "Habilitar la asignación y reasignación de fisioterapeutas en tiempo real tanto en la reserva web externa como en el panel administrativo, disminuyendo los errores de solapamiento de horarios en un 95% durante los primeros dos meses de operación."
)
add_bullet(
    "Objetivo SMART 4 (Capacitación y Adopción):",
    "Capacitar al 100% del personal administrativo y fisioterapeutas de FISIOMEDI mediante sesiones prácticas guiadas y el manual de usuario, logrando una tasa de adopción efectiva superior al 90% en la primera semana de puesta en marcha."
)

# ==================== 5. DEFINICIÓN Y ALCANCE ====================
add_heading_1("5. Definición y Alcance del Sistema")
add_p(
    "El sistema FISIOMEDI es una aplicación web fullstack que abarca tanto el canal público de atención al cliente como los procesos internos de gestión clínica y administrativa. El alcance del proyecto comprende los siguientes módulos funcionales:"
)
add_bullet(
    "Módulo 1 - Portal Web y Landing Page (/):",
    "Presentación institucional de FISIOMEDI, catálogo interactivo de servicios terapéuticos (Traumatológica, Deportiva, Neurológica, Dolor Crónico, etc.), carrusel dinámico y canales de contacto directo."
)
add_bullet(
    "Módulo 2 - Reserva de Citas en Línea (/reservar):",
    "Formulario guiado en 4 pasos: selección de servicio, elección opcional de fisioterapeuta de preferencia (o asignación automática), selector de fecha con cálculo dinámico de turnos horarios disponibles y captura de datos del paciente con validaciones en tiempo real."
)
add_bullet(
    "Módulo 3 - Portal del Paciente 'Mi Cuenta' (/mi-cuenta):",
    "Acceso privado para pacientes mediante DNI y contraseña generada. Incluye 4 pestañas especializadas: Próximas Citas (con indicación del fisioterapeuta asignado), Mis Resultados (visor y descarga de resonancias/rayos X), Mis Datos (actualización de dirección, teléfono y notas) e Historial de citas."
)
add_bullet(
    "Módulo 4 - Panel Administrativo de Citas (/admin/citas):",
    "Bandeja central de citas con filtros por estado (Pendientes, Confirmadas, Completadas, Canceladas), filtro desplegable por médico/fisioterapeuta, agendamiento de citas internas, selector interactivo de especialista en cada fila y acción de cambio de estado a un clic."
)
add_bullet(
    "Módulo 5 - Carga y Publicación de Resultados Médicos (Modal Integrado):",
    "Permite al profesional cargar archivos diagnósticos pesados (resonancias magnéticas, ecografías, radiografías, informes en PDF), redactar observaciones clínicas y autocompletar la cita de forma coordinada."
)
add_bullet(
    "Módulo 6 - Ficha Integral de Pacientes e Historias Clínicas (/admin/pacientes/[id]):",
    "Expediente digital con datos filiatorios, cronología de atenciones clínicas (motivo, diagnóstico, tratamiento, profesional a cargo), repositorio de exámenes médicos con opción de validación o rechazo, y generador de contraseñas de acceso al portal."
)
add_bullet(
    "Módulo 7 - Gestión de Usuarios y Seguridad (/admin/usuarios):",
    "Control de cuentas del personal clasificado en roles: Administrador, Terapeuta/Médico y Paciente. Protección con PBKDF2/SHA-256, salting criptográfico y cookies de sesión HTTP-only seguras."
)

doc.add_page_break()

# ==================== 6. MODELADO DE PROCESOS DE NEGOCIO (RUP) ====================
add_heading_1("6. Modelado de Procesos de Negocio (Metodología RUP)")
add_p(
    "En concordancia con las fases de Inicio y Elaboración de la metodología RUP (Rational Unified Process), se ha procedido al modelado formal de los procesos de negocio de FISIOMEDI, identificando los actores del negocio, los Casos de Uso del Negocio (CUN) y el flujo de actividades dinámicas entre los participantes."
)

add_heading_2("6.1. Actores del Negocio")
add_bullet(
    "Paciente (Actor del Negocio):",
    "Persona que requiere atención fisioterapéutica. Interactúa con el sistema para solicitar citas vía web, ingresar a su cuenta con su DNI, consultar su historial y descargar sus resultados de exámenes médicos."
)
add_bullet(
    "Recepcionista / Administrador (Actor del Negocio):",
    "Personal encargado de la recepción, atención al cliente y gestión operativa de la clínica. Registra citas presenciales o telefónicas, confirma reservas web, asigna o reasigna fisioterapeutas y crea expedientes de nuevos pacientes."
)
add_bullet(
    "Fisioterapeuta / Médico Especialista (Actor del Negocio):",
    "Profesional de la salud que brinda la sesión terapéutica y evalúa los exámenes clínicos. Consulta las citas asignadas a su nombre, ejecuta el tratamiento, sube los archivos de resonancia o rayos X con sus notas diagnósticas y valida los informes médicos."
)

add_heading_2("6.2. Casos de Uso del Negocio (CUN)")
add_p("Se han delimitado los siguientes seis Casos de Uso del Negocio principales:")
add_bullet("CUN-01: Gestión y Reserva de Citas (Web / Presencial):", "Comprende la selección de servicio, fecha, horario y verificación de disponibilidad para asentar una cita en el sistema.")
add_bullet("CUN-02: Asignación de Fisioterapeuta Especialista:", "Permite vincular a un médico o fisioterapeuta específico con la cita médica, garantizando la continuidad del tratamiento del paciente.")
add_bullet("CUN-03: Atención e Ingreso de Resultados Diagnósticos:", "Abarca la ejecución de la terapia y la carga del informe o estudio en PDF/imagen (resonancia magnética, radiografía, ecografía) con diagnóstico profesional.")
add_bullet("CUN-04: Consulta y Descarga de Resultados (Portal Paciente):", "Permite al paciente ingresar a su portal seguro para revisar sus citas y descargar directamente los archivos médicos validados.")
add_bullet("CUN-05: Mantenimiento de Fichas e Historias Clínicas:", "Administración de la información clínica, antecedentes, diagnósticos y tratamientos evolutivos del paciente.")
add_bullet("CUN-06: Control de Acceso y Gestión de Usuarios/Roles:", "Autenticación, generación de contraseñas temporales y resguardo de privilegios para administradores, terapeutas y pacientes.")

# Insert Diagram 1: Actores y CUN
diag_cun_path = os.path.join(diag_dir, "diagrama_actores_cun.png")
add_image_figure(diag_cun_path, "Modelo de Actores y Casos de Uso del Negocio (CUN) - Metodología RUP", width=Inches(6.2))

add_heading_2("6.3. Diagrama de Actividades RUP (Flujo Integral con Carriles / Swimlanes)")
add_p(
    "El diagrama de actividades con carriles de responsabilidad (Swimlanes) ilustra la interacción secuencial y coordinada entre los cuatro participantes del proceso de negocio: el Paciente, el Recepcionista, el Fisioterapeuta y el Sistema Web FisioMedi con su base de datos PostgreSQL."
)

# Insert Diagram 2: Actividades RUP
diag_act_path = os.path.join(diag_dir, "diagrama_actividades_rup.png")
add_image_figure(diag_act_path, "Diagrama de Actividades RUP: Flujo Integral de Reserva, Asignación, Atención y Publicación de Resultados", width=Inches(6.2))

add_heading_2("6.4. Especificación Detallada de los Procesos Principales")

add_heading_3("Especificación CUN-01: Gestión y Reserva de Citas")
add_bullet("Precondición:", "El sistema debe estar en línea y con turnos de atención configurados en la base de datos.")
add_bullet("Flujo Normal Web:", "1. El paciente accede a /reservar. 2. Selecciona el servicio y terapeuta preferido. 3. Elige la fecha y un turno disponible. 4. Completa sus datos (DNI, teléfono, nombre). 5. El sistema registra la cita en estado 'pendiente' y emite la confirmación.")
add_bullet("Flujo Normal Interno:", "1. El recepcionista ingresa a /admin/citas. 2. Abre '+ Agendar cita interna'. 3. Selecciona el paciente existente o digita datos nuevos. 4. Asigna terapeuta y horario. 5. El sistema registra la cita inmediatamente confirmada.")
add_bullet("Flujos Alternos:", "Si el turno deseado ya está ocupado, el sistema bloquea el botón y sugiere los turnos libres más próximos.")
add_bullet("Postcondición:", "La cita queda persistida en la tabla 'appointments' y visible en la agenda del especialista.")

add_heading_3("Especificación CUN-03: Atención e Ingreso de Resultados Médicos")
add_bullet("Precondición:", "La cita debe encontrarse en estado 'confirmada' o 'completada'.")
add_bullet("Flujo Normal:", "1. El fisioterapeuta abre /admin/citas y ubica la cita del paciente. 2. Hace clic en 'Ingresar resultados'. 3. En el modal emergente, adjunta el archivo (PDF de resonancia magnética o imagen de rayos X). 4. Redacta el diagnóstico y conclusiones clínicas. 5. Activa la casilla de autocompletado y guarda. 6. El sistema almacena el archivo de forma segura, registra el examen con estado 'validado' y actualiza la cita a 'completada'.")
add_bullet("Postcondición:", "El paciente puede visualizar y descargar el estudio inmediatamente desde su portal personal.")

doc.add_page_break()

# ==================== 7. PLAN DE CAPACITACIÓN ====================
add_heading_1("7. Plan de Capacitación a Usuarios Finales")
add_p(
    "Para garantizar la adopción exitosa y eficiente de la plataforma web por parte de todo el equipo de FISIOMEDI, se ha elaborado un programa de capacitación teórico-práctico estructurado en tres audiencias objetivo:"
)

# Table of training
plan_tbl = doc.add_table(rows=4, cols=4)
plan_tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
plan_tbl.autofit = False

headers = ["Audiencia", "Duración", "Modalidad", "Temario Principal"]
for i, h in enumerate(headers):
    c = plan_tbl.rows[0].cells[i]
    set_cell_background(c, "1e3a8a")
    p = c.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(h)
    r.font.name = "Arial"
    r.font.size = Pt(8.5)
    r.font.bold = True
    r.font.color.rgb = RGBColor(255, 255, 255)

data_plan = [
    ("Personal de Recepción y Administración", "2 sesiones de 60 min", "Presencial en clínica con equipos reales", "• Gestión del panel de citas y filtros por médico\n• Agendamiento de citas internas\n• Registro de nuevos pacientes y generación de contraseñas\n• Validación de reservas web y reprogramaciones"),
    ("Fisioterapeutas y Personal Médico", "1 sesión de 45 min", "Presencial / Práctica guiada", "• Consulta de citas y pacientes asignados\n• Subida de estudios (resonancias, ecografías, rayos X)\n• Redacción de notas diagnósticas en modal clínico\n• Registro de evoluciones en la ficha médica"),
    ("Pacientes (Inducción digital)", "Asincrónica continua", "Guía interactiva en la web y WhatsApp", "• Uso del portal de reserva web en 4 pasos\n• Ingreso con DNI y cambio de contraseña temporal\n• Descarga de resultados médicos desde el móvil")
]

for row_idx, row_data in enumerate(data_plan, start=1):
    row = plan_tbl.rows[row_idx]
    bg_color = "f8fafc" if row_idx % 2 == 1 else "ffffff"
    for col_idx, text in enumerate(row_data):
        cell = row.cells[col_idx]
        set_cell_background(cell, bg_color)
        p = cell.paragraphs[0]
        r = p.add_run(text)
        r.font.name = "Arial"
        r.font.size = Pt(8.5)
        r.font.color.rgb = COLOR_DARK

doc.add_page_break()

# ==================== 8. MANUAL DE USO PARA USUARIOS FINALES ====================
add_heading_1("8. Manual de Uso para Usuarios Finales (Guía de Interfaces)")
add_p(
    "A continuación se presenta el manual descriptivo y visual de los módulos funcionales del Sistema Web FISIOMEDI, documentando cada interfaz desarrollada y su modo de empleo."
)

add_heading_2("8.1. Módulo de Reserva de Citas en Línea (/reservar)")
add_p(
    "El módulo de reserva permite a los pacientes solicitar una cita desde cualquier dispositivo sin necesidad de llamar a la clínica. El formulario opera en 4 pasos integrados:"
)
add_bullet("Paso 1: Selección de Servicio:", "El paciente elige la especialidad requerida (ej. Fisioterapia Traumatológica, Terapia Manual, etc.). Se visualiza la tarifa y duración de la sesión.")
add_bullet("Paso 2: Especialista de Preferencia:", "Permite elegir un fisioterapeuta en específico o seleccionar 'Cualquier especialista disponible' para una asignación automática y flexible.")
add_bullet("Paso 3: Selector de Fecha y Turnos:", "Al elegir un día hábil, el sistema consulta en tiempo real a PostgreSQL los horarios libres y despliega botones con las horas disponibles.")
add_bullet("Paso 4: Datos del Paciente:", "Captura de nombre completo, teléfono y motivo de consulta. Al pulsar 'Confirmar y Reservar Cita', se genera el registro en la base de datos.")

# Image UI Reservar
add_image_figure(os.path.join(diag_dir, "ui_reservar.png"), "Interfaz Web del Formulario de Reserva de Citas (/reservar)", width=Inches(6.0))

add_heading_2("8.2. Portal del Paciente 'Mi Cuenta' (/mi-cuenta)")
add_p(
    "Los pacientes acceden mediante su DNI y contraseña generada. Este portal centraliza su experiencia con la clínica a través de cuatro secciones principales:"
)
add_bullet("Pestaña Próximas Citas:", "Muestra la fecha, hora, servicio y el nombre del fisioterapeuta asignado ('👨‍⚕️ Atendido por: Dr. Pedro Fernández').")
add_bullet("Pestaña Mis Resultados (Estudios Médicos):", "Repositorio donde el paciente puede visualizar las resonancias magnéticas, ecografías y radiografías subidas por el especialista. Cada archivo cuenta con su título, fecha, observaciones clínicas y botón 'Descargar PDF'.")
add_bullet("Pestaña Mis Datos:", "Permite al paciente actualizar su dirección de residencia, teléfono de contacto y antecedentes.")
add_bullet("Pestaña Historial de Citas:", "Registro histórico de atenciones concluidas o citas pasadas.")

# Image UI Mi Cuenta
add_image_figure(os.path.join(diag_dir, "ui_mi_cuenta.png"), "Interfaz del Portal del Paciente con Pestaña de Resultados Médicos (/mi-cuenta)", width=Inches(6.0))

add_heading_2("8.3. Panel Administrativo: Gestión de Citas y Asignación de Médicos (/admin/citas)")
add_p(
    "Esta interfaz es el centro de control del recepcionista y administrador. Sus componentes clave son:"
)
add_bullet("Filtros por Estado:", "Botones para visualizar citas Activas, Pendientes, Confirmadas, Completadas o Canceladas.")
add_bullet("Filtro Dinámico por Especialista:", "Lista desplegable que filtra instantáneamente las citas de un fisioterapeuta específico.")
add_bullet("Asignador Interactivo de Terapeuta:", "Selector en cada fila de la tabla que permite asignar o cambiar el fisioterapeuta responsable con un solo clic.")
add_bullet("Botón 'Ingresar Resultados':", "Disponible en citas confirmadas o completadas, abre el modal de subida de estudios clínicos.")

# Image UI Admin Citas
add_image_figure(os.path.join(diag_dir, "ui_admin_citas.png"), "Panel Administrativo de Citas con Asignación de Terapeuta y Modal de Resultados", width=Inches(6.0))

add_heading_2("8.4. Ficha Integral del Paciente e Historia Clínica (/admin/pacientes/[id])")
add_p(
    "Permite al personal de salud acceder al expediente completo del paciente:"
)
add_bullet("Historia Clínica Evolutiva:", "Registro cronológico de consultas, diagnósticos médicos y tratamientos fisioterapéuticos aplicados.")
add_bullet("Exámenes y Documentos Archivados:", "Gestor donde se pueden revisar, validar o rechazar los exámenes médicos subidos.")
add_bullet("Gestión de Cuenta de Acceso:", "Genera automáticamente la cuenta de usuario (usuario = DNI) y una contraseña aleatoria inicial que el paciente deberá cambiar al ingresar.")

# Image UI Ficha Paciente
add_image_figure(os.path.join(diag_dir, "ui_ficha_paciente.png"), "Ficha Clínica del Paciente con Historia Clínica y Repositorio de Exámenes", width=Inches(6.0))

doc.add_page_break()

# ==================== 9. PRODUCTOS Y ENTREGABLES ====================
add_heading_1("9. Productos y Entregables del Software")
add_p(
    "Como resultado del proyecto de software, se entregan los componentes arquitectónicos, el modelo relacional de datos y el código fuente desplegado en producción:"
)

add_heading_2("9.1. Arquitectura Tecnológica del Software")
add_p(
    "La solución adopta una arquitectura desacoplada en capas optimizada para entornos cloud serverless:"
)
add_bullet("Capa de Presentación (Frontend):", "Desarrollada con Next.js 16 y Tailwind CSS. Implementa un diseño adaptable (responsive design) compatible con celulares, tablets y computadoras.")
add_bullet("Capa de Aplicación y Negocio (Backend):", "React Server Components (RSC) para renderizado ultra veloz en servidor, y Next.js Server Actions para transacciones seguras sin exponer endpoints vulnerables.")
add_bullet("Capa de Datos (Persistencia):", "Base de datos relacional PostgreSQL alojada en Neon Serverless con pool de conexiones PgBouncer y cifrado SSL.")
add_bullet("Almacenamiento de Archivos Clínicos:", "Almacenamiento seguro de archivos pesados (PDFs de resonancias magnéticas, rayos X) con validación de tipo MIME y límite de 10 MB.")

# Image Diagram Architecture
add_image_figure(os.path.join(diag_dir, "diagrama_arquitectura_web.png"), "Diagrama de Arquitectura Tecnológica del Sistema Web FISIOMEDI", width=Inches(6.0))

add_heading_2("9.2. Modelo Relacional de Base de Datos (PostgreSQL Neon)")
add_p(
    "A diferencia del almacenamiento en archivos planos, el sistema actual cuenta con un esquema de base de datos relacional normalizado con claves foráneas, índices de búsqueda acelerada y restricciones de integridad:"
)

# Image Diagram ER
add_image_figure(os.path.join(diag_dir, "diagrama_bd_er.png"), "Diagrama Entidad-Relación de la Base de Datos (PostgreSQL Neon)", width=Inches(6.0))

# Table describing schema
db_tbl = doc.add_table(rows=6, cols=3)
db_tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
db_tbl.autofit = False

db_hdrs = ["Tabla", "Descripción Funcional", "Claves y Relaciones Clave"]
for i, h in enumerate(db_hdrs):
    c = db_tbl.rows[0].cells[i]
    set_cell_background(c, "1e3a8a")
    p = c.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(h)
    r.font.name = "Arial"
    r.font.size = Pt(8.5)
    r.font.bold = True
    r.font.color.rgb = RGBColor(255, 255, 255)

data_db = [
    ("users", "Almacena cuentas de usuarios (admin, terapeutas y pacientes) con salting y hash de contraseñas.", "PK: id (UUID)\nFK: patient_id -> patients(id)"),
    ("patients", "Registro clínico y filiatorio de los pacientes atendidos en la clínica.", "PK: id (UUID)\nCampos: doc_id (DNI), name, phone, email, notes"),
    ("appointments", "Almacena las citas agendadas tanto vía web como registradas internamente.", "PK: id (UUID)\nFK: patient_id -> patients(id)\nFK: therapist_id -> users(id)"),
    ("exams", "Repositorio de resultados diagnósticos (resonancias magnéticas, ecografías, rayos X).", "PK: id (UUID)\nFK: patient_id -> patients(id)\nFK: appointment_id -> appointments(id)"),
    ("history_entries", "Cronología de sesiones terapéuticas, motivos, diagnósticos y tratamientos evolutivos.", "PK: id (UUID)\nFK: patient_id -> patients(id)")
]

for row_idx, row_data in enumerate(data_db, start=1):
    row = db_tbl.rows[row_idx]
    bg_color = "f8fafc" if row_idx % 2 == 1 else "ffffff"
    for col_idx, text in enumerate(row_data):
        cell = row.cells[col_idx]
        set_cell_background(cell, bg_color)
        p = cell.paragraphs[0]
        r = p.add_run(text)
        r.font.name = "Arial"
        r.font.size = Pt(8.5)
        r.font.color.rgb = COLOR_DARK

doc.add_page_break()

# ==================== 10. CONCLUSIONES ====================
add_heading_1("10. Conclusiones")
add_bullet(
    "Conclusión 1 (Cumplimiento de Objetivos):",
    "Se logró con éxito el desarrollo, modelado e implementación del Sistema Web FISIOMEDI, migrando el 100% de la gestión de citas de papel a una plataforma digital moderna y automatizada."
)
add_bullet(
    "Conclusión 2 (Innovación en Resultados Médicos):",
    "La incorporación del módulo 'Ingresar resultados' y el portal 'Mi Cuenta' resuelve de manera definitiva el problema de entrega de estudios diagnósticos pesados (resonancias, ecografías y radiografías), brindando un valor diferencial frente a otros centros de fisioterapia."
)
add_bullet(
    "Conclusión 3 (Eficiencia en Asignación de Terapeutas):",
    "La funcionalidad de asignación de especialistas en tiempo real tanto en la web pública como en el panel administrativo evita solapamientos de agendas y permite un balance equitativo de carga de trabajo entre el equipo médico."
)
add_bullet(
    "Conclusión 4 (Arquitectura Tecnológica Robusta):",
    "El uso de Next.js 16 con PostgreSQL Neon demuestra que las tecnologías serverless ofrecen un rendimiento de alta velocidad (SSR) y máxima seguridad clínica sin incurrir en elevados costos de infraestructura física."
)
add_bullet(
    "Conclusión 5 (Consolidación de Competencias Profesionales):",
    "El proyecto consolidó las competencias profesionales del equipo de estudiantes en análisis RUP, diseño de arquitecturas web fullstack, gestión de bases de datos relacionales y control de versiones colaborativo con Git y GitHub."
)

# ==================== 11. RECOMENDACIONES ====================
add_heading_1("11. Recomendaciones")
add_bullet(
    "Recomendación 1 (Copias de Seguridad Automatizadas):",
    "Programar copias de seguridad automáticas diarias de la base de datos PostgreSQL en Neon hacia un bucket de almacenamiento seguro (ej. AWS S3 o Google Cloud Storage) para prevenir cualquier contingencia de pérdida de datos clínicos."
)
add_bullet(
    "Recomendación 2 (Notificaciones por WhatsApp y Correo):",
    "Integrar un servicio de mensajería automática (ej. Twilio API o Resend) que envíe un recordatorio por WhatsApp o correo electrónico al paciente 24 horas antes de su cita y una alerta cuando sus resultados de resonancia o rayos X hayan sido publicados."
)
add_bullet(
    "Recomendación 3 (Pasarela de Pagos en Línea):",
    "Incorporar una pasarela de pago digital (ej. Niubiz, MercadoPago o Yape/Plin) para que el paciente pueda abonar el costo de la sesión terapéutica al momento de reservar en la web."
)
add_bullet(
    "Recomendación 4 (Visor DICOM Integrado):",
    "En una siguiente versión, añadir un visor de imágenes médicas nativo (estándar DICOM) que permita al fisioterapeuta manipular contrastes y zooms en placas radiográficas directamente desde el navegador web."
)

# ==================== 12. GLOSARIO TÉCNICO ====================
add_heading_1("12. Glosario Técnico")
add_bullet("Next.js 16:", "Framework de desarrollo fullstack basado en React que ofrece renderizado del lado del servidor (SSR), optimización de recursos y Server Actions.")
add_bullet("React Server Components (RSC):", "Paradigma de componentes que se ejecutan exclusivamente en el servidor, reduciendo el peso de JavaScript descargado por el navegador.")
add_bullet("PostgreSQL:", "Sistema de gestión de bases de datos relacional de código abierto reconocido por su robustez, integridad referencial y soporte de tipos complejos como UUID y JSONB.")
add_bullet("Neon Serverless:", "Servicio en la nube que provee bases de datos PostgreSQL serverless con escalado automático y poolers de conexión de alta concurrencia.")
add_bullet("Metodología RUP:", "Metodología disciplinada de desarrollo de software basada en casos de uso, iteraciones y modelado visual con UML.")
add_bullet("Caso de Uso del Negocio (CUN):", "Secuencia de acciones que una empresa realiza para generar un resultado observable de valor para un actor del negocio.")
add_bullet("Diagrama Swimlane (Carriles de Natación):", "Variante del diagrama de actividades que organiza las tareas en columnas para identificar visualmente qué actor o sistema ejecuta cada paso.")
add_bullet("PBKDF2 y Salting Criptográfico:", "Algoritmo de derivación de claves que aplica miles de iteraciones de hashing con una cadena aleatoria (salt) para proteger contraseñas contra ataques de fuerza bruta.")
add_bullet("Vercel:", "Plataforma cloud global especializada en el despliegue automático, balanceo de carga y distribución perimetral de aplicaciones web modernas.")
add_bullet("Resonancia Magnética (RMN):", "Estudio de diagnóstico por imágenes de alta resolución utilizado en fisioterapia para evaluar tejidos blandos, ligamentos y discos vertebrales.")

# Save Document
doc.save(output_docx_web)
print("Saved document to web docs:", output_docx_web)

# Copy to root docs folder so it's in c:\Users\user\Documents\Fisiomedi\docs
try:
    shutil.copy2(output_docx_web, output_docx_root)
    print("Copied document to root docs:", output_docx_root)
except Exception as e:
    print("Copy error:", e)
